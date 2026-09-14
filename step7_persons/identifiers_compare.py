import pandas as pd
from paths import PATH_WORK
 
# ---------------------------------------------------------------
# 6. Comparaison ORCID post-fusion (orcid_runner vs idref)
# ---------------------------------------------------------------
def check_merge_integrity(df_left: pd.DataFrame, df_right: pd.DataFrame,
                           key_cols=("last_name", "first_name"),
                           left_label="orcid_pipeline", right_label="idref",
                           left_id_col="orcid_id_final", right_id_col="idref_id"):
    """
    À appeler AVANT le merge, sur les deux tables sources.
    Vérifie que la clé de jointure (nom+prénom) est bien unique dans
    chaque table -> sinon le merge("outer") va créer un produit
    cartésien local et dupliquer des lignes silencieusement.
 
    Un même nom+prénom peut légitimement apparaître plusieurs fois côté
    orcid_pipeline (une personne avec plusieurs projets/rôles) -> ce
    n'est PAS un risque d'homonymie tant que toutes ces lignes partagent
    le même identifiant (left_id_col / right_id_col). On distingue donc :
      - risque_homonyme=False : doublon bénin (même identifiant partout,
        ou identifiant absent -> rien à comparer)
      - risque_homonyme=True  : identifiants DIFFÉRENTS sous le même nom
        -> vraisemblablement deux personnes distinctes, à vérifier
    """
    def _key(df):
        return (df[key_cols[0]].astype(str).str.strip().str.lower() + "||" +
                df[key_cols[1]].astype(str).str.strip().str.lower())
 
    def _flag_homonyme(df, id_col):
        df = df.copy()
        key = _key(df)
        df["_key"] = key
        dupes = df[key.duplicated(keep=False)].copy()
        if len(dupes) == 0 or id_col not in dupes.columns:
            if len(dupes):
                dupes["risque_homonyme"] = False
            return dupes.drop(columns="_key", errors="ignore")
 
        nb_ids_distincts = (
            dupes[dupes[id_col].notna()]
            .groupby("_key")[id_col].nunique()
        )
        homonymes_keys = set(nb_ids_distincts[nb_ids_distincts > 1].index)
        dupes["risque_homonyme"] = dupes["_key"].isin(homonymes_keys)
        return dupes.drop(columns="_key")
 
    left_dupes = _flag_homonyme(df_left, left_id_col)
    right_dupes = _flag_homonyme(df_right, right_id_col)
 
    n_left_risk = left_dupes["risque_homonyme"].sum() if len(left_dupes) else 0
    n_right_risk = right_dupes["risque_homonyme"].sum() if len(right_dupes) else 0
 
    print(f"[merge check] {left_label}: {len(df_left)} ligne(s), "
          f"{len(left_dupes)} ligne(s) de nom en doublon "
          f"({n_left_risk} avec identifiant DIFFÉRENT -> risque réel, "
          f"{len(left_dupes) - n_left_risk} bénin(es) -> même identifiant/projets multiples)")
    print(f"[merge check] {right_label}: {len(df_right)} ligne(s), "
          f"{len(right_dupes)} ligne(s) de nom en doublon "
          f"({n_right_risk} avec identifiant DIFFÉRENT -> risque réel, "
          f"{len(right_dupes) - n_right_risk} bénin(es))")
 
    if n_left_risk:
        print(f"\n⚠️  Doublons à RISQUE côté {left_label} (identifiants différents sous le même nom) :")
        cols_show = list(key_cols) + ([left_id_col] if left_id_col in left_dupes.columns else [])
        print(left_dupes[left_dupes["risque_homonyme"]][cols_show].to_string(index=False))
    if n_right_risk:
        print(f"\n⚠️  Doublons à RISQUE côté {right_label} (identifiants différents sous le même nom) :")
        cols_show = list(key_cols) + ([right_id_col] if right_id_col in right_dupes.columns else [])
        print(right_dupes[right_dupes["risque_homonyme"]][cols_show].to_string(index=False))
 
    return left_dupes, right_dupes
 
 
def export_verification_excel(left_dupes: pd.DataFrame, right_dupes: pd.DataFrame,
                               mismatches: pd.DataFrame, output_path: str):
    """
    Exporte les 3 données à vérifier manuellement dans un classeur Excel,
    un onglet par catégorie :
      - "Doublons cle orcid"  : doublons de clé côté pipeline ORCID
      - "Doublons cle idref"  : doublons de clé côté idref
      - "Mismatch orcid"      : lignes où orcid_pipeline != orcid_idref
    """
    from openpyxl.styles import Font
 
    sheets = {
        "Doublons cle orcid": left_dupes,
        "Doublons cle idref": right_dupes,
        "Mismatch orcid": mismatches,
    }
 
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            (df if len(df) else pd.DataFrame({"info": ["aucune ligne à vérifier"]})) \
                .to_excel(writer, sheet_name=sheet_name, index=False)
 
    # Police professionnelle (Arial) sur toutes les cellules de tous les onglets
    from openpyxl import load_workbook
    wb = load_workbook(output_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", bold=(cell.row == 1))
        for col_cells in ws.columns:
            max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 60)
    wb.save(output_path)
 
    print(f"\n📄 Fichier de vérification exporté : {output_path}")
    print(f"   - Doublons cle orcid : {len(left_dupes)} ligne(s)")
    print(f"   - Doublons cle idref : {len(right_dupes)} ligne(s)")
    print(f"   - Mismatch orcid     : {len(mismatches)} ligne(s)")
 
 
def merge_orcid_idref(perso_orcid: pd.DataFrame, idref_res: pd.DataFrame,
                       key_cols=("last_name", "first_name"),
                       export_excel=True, excel_path=None):
    """
    Fusionne le résultat du pipeline ORCID (perso_orcid, avec
    orcid_id_final) et le résultat IdRef (idref_res, indépendant),
    sur nom+prénom. Vérifie l'intégrité du merge avant de le faire,
    calcule le statut de comparaison ORCID après, et exporte un fichier
    Excel de vérification (doublons + mismatches) si export_excel=True.
    """
    left_dupes, right_dupes = check_merge_integrity(
        perso_orcid, idref_res, key_cols=key_cols,
        left_label="orcid_pipeline", right_label="idref",
    )
    n_left_risk = left_dupes["risque_homonyme"].sum() if len(left_dupes) else 0
    n_right_risk = right_dupes["risque_homonyme"].sum() if len(right_dupes) else 0
    if n_left_risk or n_right_risk:
        print(f"\n⚠️  {n_left_risk + n_right_risk} doublon(s) à RISQUE réel (identifiants "
              f"différents sous le même nom) -> le merge par nom+prénom seul n'est pas fiable "
              f"pour ces personnes, vérifie-les avant d'utiliser le résultat.")
 
    def _key(df):
        return (df[key_cols[0]].astype(str).str.strip().str.lower() + "||" +
                df[key_cols[1]].astype(str).str.strip().str.lower())
 
    perso_orcid = perso_orcid.copy()
    idref_res = idref_res.copy()
    perso_orcid["_key"] = _key(perso_orcid)
    idref_res["_key"] = _key(idref_res)
 
    n_before = len(perso_orcid) + len(idref_res)
 
    cols_from_idref = [c for c in idref_res.columns
                        if c not in key_cols and c not in ("orcid_id_final", "_key")]
    merged = perso_orcid.merge(
        idref_res[["_key"] + cols_from_idref],
        on="_key", how="outer", indicator=True,
    ).drop(columns="_key")
 
    print(f"\n[merge] {len(perso_orcid)} (orcid_pipeline) + {len(idref_res)} (idref) "
          f"-> {len(merged)} ligne(s) après merge outer")
    print(merged["_merge"].value_counts().rename({
        "left_only": "orcid_pipeline seulement",
        "right_only": "idref seulement",
        "both": "présent des deux côtés",
    }).to_string())
    if len(merged) > n_before:
        print(f"\n⚠️  Le merge a AUGMENTÉ le nombre de lignes attendu -> "
              f"il y a probablement des doublons de clé non résolus.")
 
    merged = merged.rename(columns={"_merge": "merge_source"})
    result = compare_orcid_sources(merged)
 
    if export_excel:
        mismatches = result[result["orcid_comparison"] == "mismatch"]
        if excel_path is None:
            excel_path = f"{PATH_WORK}verification_merge_orcid_idref.xlsx"
        export_verification_excel(left_dupes, right_dupes, mismatches, excel_path)
 
    return result
 
 
def compare_orcid_sources(df: pd.DataFrame,
                           orcid_pipeline_col="orcid_id_final",
                           orcid_idref_col="idref_orcid_found") -> pd.DataFrame:
    """
    Ajoute une colonne orcid_comparison classifiant chaque ligne :
      - "match"              : les deux orcid concordent
      - "mismatch"            : les deux orcid sont présents mais DIFFÉRENTS
                                 -> à vérifier en priorité, indique probablement
                                 une mauvaise résolution d'un côté ou l'autre
      - "only_pipeline"       : orcid trouvé par le pipeline ORCID seulement
      - "only_idref"          : orcid trouvé via IdRef seulement (pipeline
                                 ORCID n'a rien donné, mais la notice idref
                                 contenait un orcid)
      - "idref_no_orcid"      : idref trouvé mais SANS orcid dans la notice
                                 -> pas de comparaison possible
      - "neither"             : aucun orcid nulle part
    """
    df = df.copy()
 
    a = df[orcid_pipeline_col].astype(str).str.strip()
    b = df[orcid_idref_col].astype(str).str.strip() if orcid_idref_col in df.columns else pd.Series([""] * len(df), index=df.index)
    has_a = df[orcid_pipeline_col].notna() & (a != "") & (a.str.lower() != "nan")
    has_b = df[orcid_idref_col].notna() & (b != "") & (b.str.lower() != "nan") if orcid_idref_col in df.columns else pd.Series(False, index=df.index)
    has_idref_notice = df["idref_id"].notna() if "idref_id" in df.columns else pd.Series(False, index=df.index)
 
    conditions = [
        (has_a & has_b & (a == b)),
        (has_a & has_b & (a != b)),
        (has_a & ~has_b),
        (~has_a & has_b),
        (~has_a & ~has_b & has_idref_notice),
        (~has_a & ~has_b & ~has_idref_notice),
    ]
    labels = ["match", "mismatch", "only_pipeline", "only_idref", "idref_no_orcid", "neither"]
 
    df["orcid_comparison"] = pd.Series(pd.NA, index=df.index, dtype="object")
    for cond, label in zip(conditions, labels):
        df.loc[cond, "orcid_comparison"] = label
 
    print("\n[comparaison ORCID pipeline <-> IdRef]")
    print(df["orcid_comparison"].value_counts(dropna=False).to_string())
 
    mismatches = df[df["orcid_comparison"] == "mismatch"]
    if len(mismatches):
        print(f"\n🔴 {len(mismatches)} mismatch(es) à vérifier en priorité :")
        cols_show = ["last_name", "first_name", orcid_pipeline_col, orcid_idref_col]
        cols_show = [c for c in cols_show if c in df.columns]
        print(mismatches[cols_show].to_string(index=False))
 
    only_idref = df[df["orcid_comparison"] == "only_idref"]
    if len(only_idref):
        print(f"\n🟢 {len(only_idref)} personne(s) avec un orcid trouvé SEULEMENT "
              f"via IdRef (le pipeline ORCID les avait manquées) :")
        cols_show = ["last_name", "first_name", orcid_idref_col, "idref_id"]
        cols_show = [c for c in cols_show if c in df.columns]
        print(only_idref[cols_show].to_string(index=False))
 
    return df